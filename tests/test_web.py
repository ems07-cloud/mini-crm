"""Проверки через настоящий веб-интерфейс: вход, права, страницы, формы, Excel."""
from io import BytesIO

from openpyxl import load_workbook


def test_pages_require_login(client_for):
    c = client_for(None)
    for url in ("/", "/deals", "/clients", "/tasks", "/users", "/deals/1", "/export/deals.xlsx"):
        r = c.get(url, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/login", url


def test_wrong_password(client_for):
    r = client_for(None).post("/login", data={"email": "anna@demo.ru", "password": "bad"})
    assert r.status_code == 401 and "Неверный email или пароль" in r.text


def test_all_pages_render_for_admin(client_for):
    c = client_for("boss@demo.ru")
    for url, text in [("/", "Отдел продаж"), ("/deals", "КП отправлено"), ("/clients", "СтройМонтаж"),
                      ("/clients/1", "Корпоративный сайт"), ("/deals/1", "История"), ("/tasks", "просрочена"),
                      ("/users", "Сотрудники")]:
        r = c.get(url)
        assert r.status_code == 200 and text in r.text, url


def test_manager_cannot_open_foreign_deal_or_admin_pages(client_for):
    c = client_for("anna@demo.ru")
    assert c.get("/deals/3").status_code == 404          # сделка Игоря
    assert c.get("/users").status_code == 404
    assert "Игорь Смирнов" not in c.get("/deals").text   # на доске только свои
    r = c.post("/deals/3/stage", data={"stage": "won"})
    assert r.status_code == 404


def test_deal_flow_through_forms(client_for):
    c = client_for("anna@demo.ru")
    r = c.post("/deals", data={"title": "Редизайн сайта", "client_id": 2, "amount": 150000}, follow_redirects=False)
    deal_url = r.headers["location"]
    c.post(deal_url + "/stage", data={"stage": "talks"})
    c.post(deal_url + "/note", data={"text": "Клиент просит три варианта дизайна"})
    c.post("/tasks", data={"title": "Показать варианты", "due": "2030-01-10", "deal_id": deal_url.split("/")[-1]})
    page = c.get(deal_url).text
    assert "Этап: Новая → Переговоры" in page
    assert "Клиент просит три варианта дизайна" in page
    assert "Показать варианты" in page


def test_admin_adds_manager_who_can_log_in(client_for):
    boss = client_for("boss@demo.ru")
    r = boss.post("/users", data={"name": "Олег Новиков", "email": "oleg@demo.ru", "password": "pass12345",
                                  "role": "manager"}, follow_redirects=False)
    assert r.status_code == 303
    assert "Пароль — минимум 8" in boss.post("/users", data={"name": "X", "email": "x@demo.ru",
                                                              "password": "123", "role": "manager"}).text
    oleg = client_for(None)
    assert oleg.post("/login", data={"email": "oleg@demo.ru", "password": "pass12345"},
                     follow_redirects=False).status_code == 303


def test_excel_export_respects_rights(client_for):
    def rows(email):
        r = client_for(email).get("/export/deals.xlsx")
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        return list(load_workbook(BytesIO(r.content))["Сделки"].iter_rows(min_row=2, values_only=True))

    assert len(rows("boss@demo.ru")) == 15
    anna = rows("anna@demo.ru")
    assert anna and all(r[3] == "Анна Ковалёва" for r in anna)


def test_logout(client_for):
    c = client_for("anna@demo.ru")
    c.get("/logout")
    assert c.get("/", follow_redirects=False).status_code == 303
