from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Deal

HEADERS = ["№", "Сделка", "Клиент", "Ответственный", "Этап", "Сумма, ₽", "Создана", "Закрыта"]
WIDTHS = [6, 34, 28, 20, 16, 14, 17, 17]


def deals_xlsx(deals: list[Deal]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сделки"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for d in deals:
        ws.append([d.id, d.title, d.client.name, d.owner.name, d.stage_title, d.amount, d.created_at, d.closed_at])
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = '# ##0 "₽"'
        row[6].number_format = row[7].number_format = "DD.MM.YYYY"
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if deals:
        ws.auto_filter.ref = f"A1:H{ws.max_row}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
